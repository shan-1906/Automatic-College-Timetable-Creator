"""
College Timetable Generator System
Main Flask Application
"""

import os
import json
from functools import wraps
from datetime import datetime

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, send_file, jsonify
)
from werkzeug.security import generate_password_hash, check_password_hash
import mysql.connector
from mysql.connector import Error

from config import Config

# ---------------------------------------------------------------------------
# App Initialisation
# ---------------------------------------------------------------------------
app = Flask(__name__)
app.config.from_object(Config)

# Ensure export folder exists
os.makedirs(app.config['EXPORT_FOLDER'], exist_ok=True)


# ---------------------------------------------------------------------------
# Database Helper
# ---------------------------------------------------------------------------
def get_db():
    """Return a MySQL connection using XAMPP defaults."""
    try:
        conn = mysql.connector.connect(
            host=app.config['MYSQL_HOST'],
            user=app.config['MYSQL_USER'],
            password=app.config['MYSQL_PASSWORD'],
            database=app.config['MYSQL_DB'],
            port=app.config['MYSQL_PORT']
        )
        return conn
    except Error as e:
        print(f"Database connection error: {e}")
        return None


def query_db(query, args=(), one=False, commit=False):
    """Execute a query and return results."""
    conn = get_db()
    if conn is None:
        return None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(query, args)
        if commit:
            conn.commit()
            return cursor.lastrowid
        results = cursor.fetchall()
        return (results[0] if results else None) if one else results
    except Error as e:
        print(f"Query error: {e}")
        if commit:
            conn.rollback()
        return None
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Authentication Decorator
# ---------------------------------------------------------------------------
def login_required(f):
    """Restrict routes to logged-in admins."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'admin_id' not in session:
            flash('Please login to access this page.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated


# ===================================================================
#  AUTH ROUTES
# ===================================================================
@app.route('/')
def index():
    if 'admin_id' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')

        admin = query_db('SELECT * FROM admin WHERE username = %s', (username,), one=True)

        if admin and check_password_hash(admin['password'], password):
            session['admin_id'] = admin['id']
            session['admin_username'] = admin['username']
            flash('Login successful!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Invalid username or password.', 'danger')
    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))


# ===================================================================
#  DASHBOARD
# ===================================================================
@app.route('/dashboard')
@login_required
def dashboard():
    stats = {
        'teachers': query_db('SELECT COUNT(*) AS c FROM teachers', one=True)['c'],
        'subjects': query_db('SELECT COUNT(*) AS c FROM subjects', one=True)['c'],
        'classes':  query_db('SELECT COUNT(*) AS c FROM classes',  one=True)['c'],
        'rooms':    query_db('SELECT COUNT(*) AS c FROM rooms',    one=True)['c'],
        'timetable_entries': query_db('SELECT COUNT(*) AS c FROM timetable', one=True)['c'],
    }
    return render_template('dashboard.html', stats=stats)


# ===================================================================
#  TEACHERS CRUD
# ===================================================================
@app.route('/teachers')
@login_required
def teachers():
    all_teachers = query_db('SELECT * FROM teachers ORDER BY name')
    return render_template('teachers.html', teachers=all_teachers)


@app.route('/teachers/add', methods=['POST'])
@login_required
def add_teacher():
    name       = request.form.get('name', '').strip()
    department = request.form.get('department', '').strip()
    max_hours  = request.form.get('max_hours', 20, type=int)

    # Build availability JSON from checkboxes
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    availability = {}
    for day in days:
        periods = request.form.getlist(f'avail_{day}')
        if periods:
            availability[day] = [int(p) for p in periods]

    if not name or not department:
        flash('Name and department are required.', 'danger')
        return redirect(url_for('teachers'))

    query_db(
        'INSERT INTO teachers (name, department, availability, max_hours) VALUES (%s, %s, %s, %s)',
        (name, department, json.dumps(availability), max_hours), commit=True
    )
    flash(f'Teacher "{name}" added successfully.', 'success')
    return redirect(url_for('teachers'))


@app.route('/teachers/edit/<int:tid>', methods=['POST'])
@login_required
def edit_teacher(tid):
    name       = request.form.get('name', '').strip()
    department = request.form.get('department', '').strip()
    max_hours  = request.form.get('max_hours', 20, type=int)

    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    availability = {}
    for day in days:
        periods = request.form.getlist(f'avail_{day}')
        if periods:
            availability[day] = [int(p) for p in periods]

    query_db(
        'UPDATE teachers SET name=%s, department=%s, availability=%s, max_hours=%s WHERE id=%s',
        (name, department, json.dumps(availability), max_hours, tid), commit=True
    )
    flash(f'Teacher "{name}" updated.', 'success')
    return redirect(url_for('teachers'))


@app.route('/teachers/delete/<int:tid>')
@login_required
def delete_teacher(tid):
    query_db('DELETE FROM teachers WHERE id=%s', (tid,), commit=True)
    flash('Teacher deleted.', 'success')
    return redirect(url_for('teachers'))


# ===================================================================
#  SUBJECTS CRUD
# ===================================================================
@app.route('/subjects')
@login_required
def subjects():
    all_subjects = query_db('''
        SELECT s.*, t.name AS teacher_name
        FROM subjects s
        LEFT JOIN teachers t ON s.teacher_id = t.id
        ORDER BY s.name
    ''')
    all_teachers = query_db('SELECT id, name FROM teachers ORDER BY name')
    return render_template('subjects.html', subjects=all_subjects, teachers=all_teachers)


@app.route('/subjects/add', methods=['POST'])
@login_required
def add_subject():
    name         = request.form.get('name', '').strip()
    weekly_hours = request.form.get('weekly_hours', 3, type=int)
    teacher_id   = request.form.get('teacher_id', type=int)
    subject_type = request.form.get('subject_type', 'theory')

    if not name:
        flash('Subject name is required.', 'danger')
        return redirect(url_for('subjects'))

    query_db(
        'INSERT INTO subjects (name, weekly_hours, teacher_id, subject_type) VALUES (%s, %s, %s, %s)',
        (name, weekly_hours, teacher_id, subject_type), commit=True
    )
    flash(f'Subject "{name}" added.', 'success')
    return redirect(url_for('subjects'))


@app.route('/subjects/edit/<int:sid>', methods=['POST'])
@login_required
def edit_subject(sid):
    name         = request.form.get('name', '').strip()
    weekly_hours = request.form.get('weekly_hours', 3, type=int)
    teacher_id   = request.form.get('teacher_id', type=int)
    subject_type = request.form.get('subject_type', 'theory')

    query_db(
        'UPDATE subjects SET name=%s, weekly_hours=%s, teacher_id=%s, subject_type=%s WHERE id=%s',
        (name, weekly_hours, teacher_id, subject_type, sid), commit=True
    )
    flash(f'Subject "{name}" updated.', 'success')
    return redirect(url_for('subjects'))


@app.route('/subjects/delete/<int:sid>')
@login_required
def delete_subject(sid):
    query_db('DELETE FROM subjects WHERE id=%s', (sid,), commit=True)
    flash('Subject deleted.', 'success')
    return redirect(url_for('subjects'))


# ===================================================================
#  CLASSES CRUD
# ===================================================================
@app.route('/classes')
@login_required
def classes():
    all_classes = query_db('SELECT * FROM classes ORDER BY department, semester, section')
    return render_template('classes.html', classes=all_classes)


@app.route('/classes/add', methods=['POST'])
@login_required
def add_class():
    department = request.form.get('department', '').strip()
    semester   = request.form.get('semester', 1, type=int)
    section    = request.form.get('section', 'A').strip()

    if not department:
        flash('Department is required.', 'danger')
        return redirect(url_for('classes'))

    query_db(
        'INSERT INTO classes (department, semester, section) VALUES (%s, %s, %s)',
        (department, semester, section), commit=True
    )
    flash('Class added successfully.', 'success')
    return redirect(url_for('classes'))


@app.route('/classes/edit/<int:cid>', methods=['POST'])
@login_required
def edit_class(cid):
    department = request.form.get('department', '').strip()
    semester   = request.form.get('semester', 1, type=int)
    section    = request.form.get('section', 'A').strip()

    query_db(
        'UPDATE classes SET department=%s, semester=%s, section=%s WHERE id=%s',
        (department, semester, section, cid), commit=True
    )
    flash('Class updated.', 'success')
    return redirect(url_for('classes'))


@app.route('/classes/delete/<int:cid>')
@login_required
def delete_class(cid):
    query_db('DELETE FROM classes WHERE id=%s', (cid,), commit=True)
    flash('Class deleted.', 'success')
    return redirect(url_for('classes'))


# ===================================================================
#  ROOMS CRUD
# ===================================================================
@app.route('/rooms')
@login_required
def rooms():
    all_rooms = query_db('SELECT * FROM rooms ORDER BY room_number')
    return render_template('rooms.html', rooms=all_rooms)


@app.route('/rooms/add', methods=['POST'])
@login_required
def add_room():
    room_number = request.form.get('room_number', '').strip()
    capacity    = request.form.get('capacity', 60, type=int)
    room_type   = request.form.get('room_type', 'classroom')

    if not room_number:
        flash('Room number is required.', 'danger')
        return redirect(url_for('rooms'))

    query_db(
        'INSERT INTO rooms (room_number, capacity, room_type) VALUES (%s, %s, %s)',
        (room_number, capacity, room_type), commit=True
    )
    flash(f'Room "{room_number}" added.', 'success')
    return redirect(url_for('rooms'))


@app.route('/rooms/edit/<int:rid>', methods=['POST'])
@login_required
def edit_room(rid):
    room_number = request.form.get('room_number', '').strip()
    capacity    = request.form.get('capacity', 60, type=int)
    room_type   = request.form.get('room_type', 'classroom')

    query_db(
        'UPDATE rooms SET room_number=%s, capacity=%s, room_type=%s WHERE id=%s',
        (room_number, capacity, room_type, rid), commit=True
    )
    flash(f'Room "{room_number}" updated.', 'success')
    return redirect(url_for('rooms'))


@app.route('/rooms/delete/<int:rid>')
@login_required
def delete_room(rid):
    query_db('DELETE FROM rooms WHERE id=%s', (rid,), commit=True)
    flash('Room deleted.', 'success')
    return redirect(url_for('rooms'))


# ===================================================================
#  TIMESLOTS CRUD
# ===================================================================
@app.route('/timeslots')
@login_required
def timeslots():
    all_slots = query_db('SELECT * FROM timeslots ORDER BY FIELD(day,"Monday","Tuesday","Wednesday","Thursday","Friday"), period_num')
    return render_template('timeslots.html', timeslots=all_slots)


@app.route('/timeslots/add', methods=['POST'])
@login_required
def add_timeslot():
    day        = request.form.get('day', '').strip()
    period_num = request.form.get('period_num', 1, type=int)
    start_time = request.form.get('start_time', '').strip()
    end_time   = request.form.get('end_time', '').strip()
    is_break   = 1 if request.form.get('is_break') else 0

    if not day or not start_time or not end_time:
        flash('All fields are required.', 'danger')
        return redirect(url_for('timeslots'))

    query_db(
        'INSERT INTO timeslots (day, period_num, start_time, end_time, is_break) VALUES (%s, %s, %s, %s, %s)',
        (day, period_num, start_time, end_time, is_break), commit=True
    )
    flash('Timeslot added.', 'success')
    return redirect(url_for('timeslots'))


@app.route('/timeslots/delete/<int:tsid>')
@login_required
def delete_timeslot(tsid):
    query_db('DELETE FROM timeslots WHERE id=%s', (tsid,), commit=True)
    flash('Timeslot deleted.', 'success')
    return redirect(url_for('timeslots'))


# ===================================================================
#  RULES
# ===================================================================
@app.route('/rules', methods=['GET', 'POST'])
@login_required
def rules():
    if request.method == 'POST':
        working_days    = request.form.get('working_days', 5, type=int)
        periods_per_day = request.form.get('periods_per_day', 6, type=int)
        break_after     = request.form.get('break_after', 3, type=int)
        break_duration  = request.form.get('break_duration', 15, type=int)
        lab_periods     = request.form.get('lab_periods', 2, type=int)

        existing = query_db('SELECT id FROM rules LIMIT 1', one=True)
        if existing:
            query_db(
                'UPDATE rules SET working_days=%s, periods_per_day=%s, break_after=%s, break_duration=%s, lab_periods=%s WHERE id=%s',
                (working_days, periods_per_day, break_after, break_duration, lab_periods, existing['id']),
                commit=True
            )
        else:
            query_db(
                'INSERT INTO rules (working_days, periods_per_day, break_after, break_duration, lab_periods) VALUES (%s,%s,%s,%s,%s)',
                (working_days, periods_per_day, break_after, break_duration, lab_periods),
                commit=True
            )
        flash('Rules updated successfully.', 'success')
        return redirect(url_for('rules'))

    current_rules = query_db('SELECT * FROM rules LIMIT 1', one=True)
    return render_template('rules.html', rules=current_rules)


# ===================================================================
#  TIMETABLE GENERATION ENGINE
# ===================================================================
@app.route('/generate', methods=['GET', 'POST'])
@login_required
def generate():
    if request.method == 'POST':
        result = generate_timetable()
        return render_template('generate.html', result=result)
    return render_template('generate.html', result=None)


def generate_timetable():
    """
    Rule-based timetable generation algorithm.

    Steps:
    1. Clear existing timetable.
    2. For each class, iterate through its subjects.
    3. For each subject, find available timeslots where:
       - Teacher is available (not already booked, matches availability)
       - Room is available (not already booked, lab→lab room, theory→classroom)
       - Weekly hours for the subject are not exceeded.
    4. Lab subjects get consecutive periods.
    5. Return a report of successful assignments and any clashes/issues.
    """
    report = {'assigned': 0, 'skipped': 0, 'clashes': [], 'details': []}

    # Clear existing timetable
    query_db('DELETE FROM timetable', commit=True)

    # Fetch all data
    all_classes  = query_db('SELECT * FROM classes ORDER BY id')
    all_subjects = query_db('SELECT * FROM subjects ORDER BY id')
    all_rooms    = query_db('SELECT * FROM rooms ORDER BY id')
    timeslots_list = query_db(
        'SELECT * FROM timeslots WHERE is_break = 0 ORDER BY FIELD(day,"Monday","Tuesday","Wednesday","Thursday","Friday"), period_num'
    )
    rules_config = query_db('SELECT * FROM rules LIMIT 1', one=True)

    if not all_classes or not all_subjects or not timeslots_list:
        report['clashes'].append('Insufficient data: please add classes, subjects, and timeslots first.')
        return report

    lab_periods = rules_config['lab_periods'] if rules_config else 2

    # Track assignments: sets for fast clash lookup
    teacher_booked = set()   # (teacher_id, timeslot_id)
    room_booked    = set()   # (room_id, timeslot_id)
    class_booked   = set()   # (class_id, timeslot_id)
    subject_hours  = {}      # (class_id, subject_id) -> count of assigned hours

    # Group timeslots by day for lab consecutive detection
    slots_by_day = {}
    for ts in timeslots_list:
        day = ts['day']
        if day not in slots_by_day:
            slots_by_day[day] = []
        slots_by_day[day].append(ts)

    # Sort each day by period number
    for day in slots_by_day:
        slots_by_day[day].sort(key=lambda x: x['period_num'])

    # ---- MAIN SCHEDULING LOOP ----
    for cls in all_classes:
        class_id = cls['id']
        class_label = f"{cls['department']} Sem-{cls['semester']} Sec-{cls['section']}"

        # Get subjects relevant to this class's department
        dept_subjects = [s for s in all_subjects if True]  # All subjects available to all classes for simplicity

        for subj in dept_subjects:
            subject_id  = subj['id']
            teacher_id  = subj['teacher_id']
            weekly_hrs  = subj['weekly_hours']
            subj_type   = subj['subject_type']

            if not teacher_id:
                report['clashes'].append(f'Subject "{subj["name"]}" has no teacher assigned — skipped.')
                report['skipped'] += 1
                continue

            # Parse teacher availability
            teacher = query_db('SELECT * FROM teachers WHERE id=%s', (teacher_id,), one=True)
            if not teacher:
                report['clashes'].append(f'Teacher ID {teacher_id} not found — skipped.')
                report['skipped'] += 1
                continue

            try:
                teacher_avail = json.loads(teacher['availability']) if teacher['availability'] else {}
            except (json.JSONDecodeError, TypeError):
                teacher_avail = {}

            key = (class_id, subject_id)
            assigned_count = subject_hours.get(key, 0)

            if subj_type == 'lab':
                # Lab: assign consecutive periods
                for day_name, day_slots in slots_by_day.items():
                    if assigned_count >= weekly_hrs:
                        break

                    avail_periods = teacher_avail.get(day_name, [])

                    for i in range(len(day_slots) - lab_periods + 1):
                        if assigned_count >= weekly_hrs:
                            break

                        consecutive = day_slots[i:i + lab_periods]

                        # Check all consecutive slots are free
                        all_free = True
                        for slot in consecutive:
                            ts_id = slot['id']
                            p_num = slot['period_num']
                            if (teacher_id, ts_id) in teacher_booked:
                                all_free = False
                                break
                            if (class_id, ts_id) in class_booked:
                                all_free = False
                                break
                            if avail_periods and p_num not in avail_periods:
                                all_free = False
                                break

                        if not all_free:
                            continue

                        # Find a lab room
                        lab_room = None
                        for room in all_rooms:
                            if room['room_type'] == 'lab':
                                room_free = all(
                                    (room['id'], slot['id']) not in room_booked
                                    for slot in consecutive
                                )
                                if room_free:
                                    lab_room = room
                                    break

                        if not lab_room:
                            continue

                        # Assign all consecutive periods
                        for slot in consecutive:
                            ts_id = slot['id']
                            query_db(
                                'INSERT INTO timetable (class_id, subject_id, teacher_id, room_id, timeslot_id) VALUES (%s,%s,%s,%s,%s)',
                                (class_id, subject_id, teacher_id, lab_room['id'], ts_id),
                                commit=True
                            )
                            teacher_booked.add((teacher_id, ts_id))
                            room_booked.add((lab_room['id'], ts_id))
                            class_booked.add((class_id, ts_id))
                            assigned_count += 1

                        subject_hours[key] = assigned_count
                        report['assigned'] += lab_periods
                        report['details'].append(
                            f'✓ {class_label} — {subj["name"]} (Lab) — {day_name} Periods {consecutive[0]["period_num"]}-{consecutive[-1]["period_num"]} — Room {lab_room["room_number"]}'
                        )
            else:
                # Theory: assign one period at a time
                for ts in timeslots_list:
                    if assigned_count >= weekly_hrs:
                        break

                    ts_id    = ts['id']
                    day_name = ts['day']
                    p_num    = ts['period_num']

                    # Check teacher availability
                    avail_periods = teacher_avail.get(day_name, [])
                    if avail_periods and p_num not in avail_periods:
                        continue

                    # Check clashes
                    if (teacher_id, ts_id) in teacher_booked:
                        continue
                    if (class_id, ts_id) in class_booked:
                        continue

                    # Find a classroom
                    assigned_room = None
                    for room in all_rooms:
                        if room['room_type'] == 'classroom':
                            if (room['id'], ts_id) not in room_booked:
                                assigned_room = room
                                break

                    if not assigned_room:
                        continue

                    # Assign
                    query_db(
                        'INSERT INTO timetable (class_id, subject_id, teacher_id, room_id, timeslot_id) VALUES (%s,%s,%s,%s,%s)',
                        (class_id, subject_id, teacher_id, assigned_room['id'], ts_id),
                        commit=True
                    )
                    teacher_booked.add((teacher_id, ts_id))
                    room_booked.add((assigned_room['id'], ts_id))
                    class_booked.add((class_id, ts_id))
                    assigned_count += 1
                    subject_hours[key] = assigned_count
                    report['assigned'] += 1
                    report['details'].append(
                        f'✓ {class_label} — {subj["name"]} — {day_name} Period {p_num} — Room {assigned_room["room_number"]}'
                    )

            # Check if all hours were assigned
            if subject_hours.get(key, 0) < weekly_hrs:
                deficit = weekly_hrs - subject_hours.get(key, 0)
                report['clashes'].append(
                    f'⚠ {class_label} — {subj["name"]}: could not assign {deficit} of {weekly_hrs} hours (insufficient slots/rooms).'
                )

    return report


# ===================================================================
#  VIEW TIMETABLE
# ===================================================================
@app.route('/view_timetable')
@login_required
def view_timetable():
    view_type = request.args.get('view', 'class')
    view_id   = request.args.get('id', type=int)
    view_day  = request.args.get('day', '')

    # Fetch filter options
    all_classes  = query_db('SELECT * FROM classes ORDER BY department, semester, section')
    all_teachers = query_db('SELECT * FROM teachers ORDER BY name')
    all_rooms    = query_db('SELECT * FROM rooms ORDER BY room_number')
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']

    timetable_data = []
    title = "Timetable"

    if view_id or view_day:
        base_query = '''
            SELECT tt.*, s.name AS subject_name, s.subject_type,
                   t.name AS teacher_name, r.room_number, r.room_type,
                   ts.day, ts.period_num, ts.start_time, ts.end_time,
                   c.department, c.semester, c.section
            FROM timetable tt
            JOIN subjects s  ON tt.subject_id  = s.id
            JOIN teachers t  ON tt.teacher_id  = t.id
            JOIN rooms r     ON tt.room_id     = r.id
            JOIN timeslots ts ON tt.timeslot_id = ts.id
            JOIN classes c   ON tt.class_id    = c.id
        '''
        conditions = []
        params = []

        if view_type == 'class' and view_id:
            conditions.append('tt.class_id = %s')
            params.append(view_id)
            cls = query_db('SELECT * FROM classes WHERE id=%s', (view_id,), one=True)
            if cls:
                title = f"{cls['department']} Sem-{cls['semester']} Sec-{cls['section']}"

        elif view_type == 'teacher' and view_id:
            conditions.append('tt.teacher_id = %s')
            params.append(view_id)
            tch = query_db('SELECT * FROM teachers WHERE id=%s', (view_id,), one=True)
            if tch:
                title = f"Teacher: {tch['name']}"

        elif view_type == 'room' and view_id:
            conditions.append('tt.room_id = %s')
            params.append(view_id)
            rm = query_db('SELECT * FROM rooms WHERE id=%s', (view_id,), one=True)
            if rm:
                title = f"Room: {rm['room_number']}"

        elif view_type == 'day' and view_day:
            conditions.append('ts.day = %s')
            params.append(view_day)
            title = f"Day: {view_day}"

        if conditions:
            base_query += ' WHERE ' + ' AND '.join(conditions)

        base_query += ' ORDER BY ts.day, ts.period_num'
        timetable_data = query_db(base_query, tuple(params))

    # Organise into a grid: days × periods
    grid = {}
    periods_set = set()
    days_in_data = set()

    if timetable_data:
        for entry in timetable_data:
            day = entry['day']
            period = entry['period_num']
            days_in_data.add(day)
            periods_set.add(period)
            key = (day, period)
            if key not in grid:
                grid[key] = []
            grid[key].append(entry)

    sorted_periods = sorted(periods_set)
    sorted_days = [d for d in days if d in days_in_data]

    return render_template(
        'view_timetable.html',
        view_type=view_type, view_id=view_id, view_day=view_day,
        classes=all_classes, teachers=all_teachers, rooms=all_rooms,
        days=days, title=title,
        grid=grid, sorted_periods=sorted_periods, sorted_days=sorted_days,
        timetable_data=timetable_data
    )


# ===================================================================
#  EXPORT ROUTES
# ===================================================================
@app.route('/export/<fmt>')
@login_required
def export_timetable(fmt):
    view_type = request.args.get('view', 'class')
    view_id   = request.args.get('id', type=int)
    view_day  = request.args.get('day', '')

    # Build query
    base_query = '''
        SELECT tt.*, s.name AS subject_name, s.subject_type,
               t.name AS teacher_name, r.room_number,
               ts.day, ts.period_num, ts.start_time, ts.end_time,
               c.department, c.semester, c.section
        FROM timetable tt
        JOIN subjects s  ON tt.subject_id  = s.id
        JOIN teachers t  ON tt.teacher_id  = t.id
        JOIN rooms r     ON tt.room_id     = r.id
        JOIN timeslots ts ON tt.timeslot_id = ts.id
        JOIN classes c   ON tt.class_id    = c.id
    '''
    conditions, params = [], []
    title = "Timetable"

    if view_type == 'class' and view_id:
        conditions.append('tt.class_id = %s')
        params.append(view_id)
        cls = query_db('SELECT * FROM classes WHERE id=%s', (view_id,), one=True)
        if cls:
            title = f"{cls['department']} Sem-{cls['semester']} Sec-{cls['section']}"
    elif view_type == 'teacher' and view_id:
        conditions.append('tt.teacher_id = %s')
        params.append(view_id)
        tch = query_db('SELECT * FROM teachers WHERE id=%s', (view_id,), one=True)
        if tch:
            title = f"Teacher - {tch['name']}"
    elif view_type == 'room' and view_id:
        conditions.append('tt.room_id = %s')
        params.append(view_id)
        rm = query_db('SELECT * FROM rooms WHERE id=%s', (view_id,), one=True)
        if rm:
            title = f"Room - {rm['room_number']}"
    elif view_type == 'day' and view_day:
        conditions.append('ts.day = %s')
        params.append(view_day)
        title = f"Day - {view_day}"

    if conditions:
        base_query += ' WHERE ' + ' AND '.join(conditions)
    base_query += ' ORDER BY ts.day, ts.period_num'

    data = query_db(base_query, tuple(params))
    if not data:
        flash('No timetable data to export.', 'warning')
        return redirect(url_for('view_timetable'))

    # Build grid
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    periods_set = sorted(set(d['period_num'] for d in data))
    days_in_data = [d for d in days if d in set(row['day'] for row in data)]
    grid = {}
    for entry in data:
        key = (entry['day'], entry['period_num'])
        if key not in grid:
            grid[key] = entry
        # Take first entry for grid cell

    if fmt == 'pdf':
        return export_pdf(title, grid, days_in_data, periods_set)
    elif fmt == 'excel':
        return export_excel(title, grid, days_in_data, periods_set)
    elif fmt == 'image':
        return export_image(title, grid, days_in_data, periods_set)
    else:
        flash('Invalid export format.', 'danger')
        return redirect(url_for('view_timetable'))


def export_pdf(title, grid, days, periods):
    """Generate a PDF timetable."""
    from fpdf import FPDF

    pdf = FPDF('L', 'mm', 'A4')
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 16)
    pdf.cell(0, 12, f'College Timetable - {title}', ln=True, align='C')
    pdf.ln(5)

    # Table header
    col_w = 45
    first_col = 25
    pdf.set_font('Helvetica', 'B', 9)
    pdf.cell(first_col, 10, 'Day', 1, 0, 'C')
    for p in periods:
        pdf.cell(col_w, 10, f'Period {p}', 1, 0, 'C')
    pdf.ln()

    # Table body
    pdf.set_font('Helvetica', '', 8)
    for day in days:
        pdf.cell(first_col, 18, day, 1, 0, 'C')
        for p in periods:
            entry = grid.get((day, p))
            if entry:
                text = f"{entry['subject_name']}\n{entry['teacher_name']}\n{entry['room_number']}"
            else:
                text = '-'
            x, y = pdf.get_x(), pdf.get_y()
            pdf.multi_cell(col_w, 6, text, 1, 'C')
            pdf.set_xy(x + col_w, y)
        pdf.ln(18)

    filepath = os.path.join(app.config['EXPORT_FOLDER'], 'timetable.pdf')
    pdf.output(filepath)
    return send_file(filepath, as_attachment=True, download_name=f'{title}.pdf')


def export_excel(title, grid, days, periods):
    """Generate an Excel timetable."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Timetable'

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(periods) + 1)
    ws.cell(1, 1, f'College Timetable - {title}').font = Font(bold=True, size=14)

    # Header row
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_font = Font(bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    ws.cell(3, 1, 'Day').font = header_font
    ws.cell(3, 1).fill = header_fill
    ws.cell(3, 1).border = thin_border
    ws.column_dimensions['A'].width = 15

    for i, p in enumerate(periods, 2):
        cell = ws.cell(3, i, f'Period {p}')
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 25

    # Data rows
    for row_idx, day in enumerate(days, 4):
        ws.cell(row_idx, 1, day).font = Font(bold=True)
        ws.cell(row_idx, 1).border = thin_border
        for col_idx, p in enumerate(periods, 2):
            entry = grid.get((day, p))
            if entry:
                text = f"{entry['subject_name']} | {entry['teacher_name']} | {entry['room_number']}"
            else:
                text = '-'
            cell = ws.cell(row_idx, col_idx, text)
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = thin_border

    filepath = os.path.join(app.config['EXPORT_FOLDER'], 'timetable.xlsx')
    wb.save(filepath)
    return send_file(filepath, as_attachment=True, download_name=f'{title}.xlsx')


def export_image(title, grid, days, periods):
    """Generate a PNG image of the timetable."""
    from PIL import Image, ImageDraw, ImageFont

    cell_w, cell_h = 180, 60
    first_col_w = 120
    header_h = 50
    title_h = 50
    width = first_col_w + cell_w * len(periods) + 20
    height = title_h + header_h + cell_h * len(days) + 20

    img = Image.new('RGB', (width, height), '#ffffff')
    draw = ImageDraw.Draw(img)

    try:
        font_title = ImageFont.truetype("arial.ttf", 20)
        font_header = ImageFont.truetype("arial.ttf", 12)
        font_cell = ImageFont.truetype("arial.ttf", 10)
    except OSError:
        font_title = ImageFont.load_default()
        font_header = font_title
        font_cell = font_title

    # Title
    draw.text((10, 10), f'College Timetable - {title}', fill='#1a1a2e', font=font_title)

    # Header
    y_start = title_h
    draw.rectangle([10, y_start, first_col_w + 10, y_start + header_h], fill='#4472C4')
    draw.text((20, y_start + 15), 'Day', fill='#ffffff', font=font_header)

    for i, p in enumerate(periods):
        x = first_col_w + 10 + i * cell_w
        draw.rectangle([x, y_start, x + cell_w, y_start + header_h], fill='#4472C4')
        draw.text((x + 5, y_start + 15), f'Period {p}', fill='#ffffff', font=font_header)

    # Data
    y_start = title_h + header_h
    for row_idx, day in enumerate(days):
        y = y_start + row_idx * cell_h
        draw.rectangle([10, y, first_col_w + 10, y + cell_h], fill='#e8eaf6', outline='#cccccc')
        draw.text((20, y + 20), day, fill='#1a1a2e', font=font_header)

        for col_idx, p in enumerate(periods):
            x = first_col_w + 10 + col_idx * cell_w
            entry = grid.get((day, p))
            if entry:
                text = f"{entry['subject_name']}\n{entry['teacher_name']}\n{entry['room_number']}"
                bg = '#f5f5f5'
            else:
                text = '-'
                bg = '#ffffff'
            draw.rectangle([x, y, x + cell_w, y + cell_h], fill=bg, outline='#cccccc')
            draw.text((x + 5, y + 5), text, fill='#333333', font=font_cell)

    filepath = os.path.join(app.config['EXPORT_FOLDER'], 'timetable.png')
    img.save(filepath)
    return send_file(filepath, as_attachment=True, download_name=f'{title}.png')


# ===================================================================
#  CLASH DETECTION API
# ===================================================================
@app.route('/api/check_clashes')
@login_required
def check_clashes():
    """Check for any scheduling clashes in the current timetable."""
    clashes = []

    # Teacher clashes
    teacher_clash = query_db('''
        SELECT t.name, ts.day, ts.period_num, COUNT(*) AS cnt
        FROM timetable tt
        JOIN teachers t ON tt.teacher_id = t.id
        JOIN timeslots ts ON tt.timeslot_id = ts.id
        GROUP BY tt.teacher_id, tt.timeslot_id
        HAVING cnt > 1
    ''')
    if teacher_clash:
        for c in teacher_clash:
            clashes.append(f"Teacher Clash: {c['name']} on {c['day']} Period {c['period_num']}")

    # Room clashes
    room_clash = query_db('''
        SELECT r.room_number, ts.day, ts.period_num, COUNT(*) AS cnt
        FROM timetable tt
        JOIN rooms r ON tt.room_id = r.id
        JOIN timeslots ts ON tt.timeslot_id = ts.id
        GROUP BY tt.room_id, tt.timeslot_id
        HAVING cnt > 1
    ''')
    if room_clash:
        for c in room_clash:
            clashes.append(f"Room Clash: {c['room_number']} on {c['day']} Period {c['period_num']}")

    # Subject overload
    overload = query_db('''
        SELECT s.name, s.weekly_hours, c.department, c.semester, c.section, COUNT(*) AS assigned
        FROM timetable tt
        JOIN subjects s ON tt.subject_id = s.id
        JOIN classes c ON tt.class_id = c.id
        GROUP BY tt.class_id, tt.subject_id
        HAVING assigned > s.weekly_hours
    ''')
    if overload:
        for o in overload:
            clashes.append(
                f"Subject Overload: {o['name']} for {o['department']} Sem-{o['semester']} "
                f"Sec-{o['section']} ({o['assigned']}/{o['weekly_hours']} hrs)"
            )

    return jsonify({'clashes': clashes, 'count': len(clashes)})


# ===================================================================
#  ADMIN PASSWORD SETUP HELPER
# ===================================================================
@app.route('/setup_admin')
def setup_admin():
    """One-time route to create/reset admin account. Remove in production."""
    hashed = generate_password_hash('admin123')
    existing = query_db('SELECT id FROM admin WHERE username=%s', ('admin',), one=True)
    if existing:
        query_db('UPDATE admin SET password=%s WHERE username=%s', (hashed, 'admin'), commit=True)
    else:
        query_db('INSERT INTO admin (username, password) VALUES (%s, %s)', ('admin', hashed), commit=True)
    flash('Admin account ready. Username: admin / Password: admin123', 'success')
    return redirect(url_for('login'))


# ===================================================================
#  RUN
# ===================================================================
if __name__ == '__main__':
    app.run(debug=True, port=5000)
